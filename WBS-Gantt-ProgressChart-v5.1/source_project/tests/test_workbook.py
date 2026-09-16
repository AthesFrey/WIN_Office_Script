"""File-level regression tests. Excel COM smoke tests are kept separately."""
import datetime as dt
import io
import posixpath
import random
import re
import sys
import warnings
from pathlib import Path
from zipfile import ZipFile

import olefile
import openpyxl
import pytest
from lxml import etree as ET
from oletools.olevba import VBA_Parser, VBA_Project, decompress_stream

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from build_workbook import EXAMPLES, NS, build
from vba_project import END, FREE, cfb_write, compress_vba


@pytest.fixture(scope="session")
def artifact(tmp_path_factory):
    return build(ROOT / "WBS.xlsm", tmp_path_factory.mktemp("release") / "WBS.xlsm")


@pytest.fixture(scope="session")
def parts(artifact):
    with ZipFile(artifact) as z:
        assert z.testzip() is None
        return {n: z.read(n) for n in z.namelist()}


def tree_signature(node):
    return node.tag, dict(node.attrib), node.text, [tree_signature(c) for c in node]


def test_two_sheets_and_no_obsolete_parts(parts):
    workbook = ET.fromstring(parts["xl/workbook.xml"])
    assert [s.get("name") for s in workbook.find("s:sheets", NS)] == ["WBS", "Milestones"]
    assert not any(n in parts for n in ["xl/worksheets/sheet2.xml", "xl/worksheets/sheet3.xml", "xl/calcChain.xml"])
    for name, data in parts.items():
        if name.endswith((".xml", ".rels")):
            assert not re.search(rb"(?:Data|Resource)!|calcChain|worksheets/sheet[23]\.xml", data), name
    titles = ET.fromstring(parts["docProps/app.xml"])
    assert titles.xpath('//*[local-name()="TitlesOfParts"]/*/*/text()') == ["WBS", "Milestones"]


def test_english_visible_content(parts):
    for name, data in parts.items():
        if name.endswith((".xml", ".vml")):
            root = ET.fromstring(data)
            assert not re.search("[\u3400-\u9fff]", " ".join(root.itertext())), name


def test_all_package_relationship_targets_exist(parts):
    for name, data in parts.items():
        if name.endswith(".xml") or name.endswith(".rels"):
            ET.fromstring(data)
        if not name.endswith(".rels"):
            continue
        root = ET.fromstring(data)
        ids = [r.get("Id") for r in root]
        assert len(ids) == len(set(ids))
        parent = posixpath.dirname(posixpath.dirname(name))
        for rel in root:
            if rel.get("TargetMode") == "External":
                continue
            target = rel.get("Target")
            full = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(parent, target))
            assert full in parts, (name, target)


def test_preserve_original_styles(parts):
    with ZipFile(ROOT / "WBS.xlsm") as z:
        before = ET.fromstring(z.read("xl/styles.xml"))
    after = ET.fromstring(parts["xl/styles.xml"])
    for group in before:
        current = after.find(group.tag)
        assert current is not None
        for i, element in enumerate(group):
            assert tree_signature(element) == tree_signature(current[i])


def test_v5_vba_modules_embedded(artifact, parts):
    parser = VBA_Parser(str(artifact))
    try:
        sources = {name: code.replace("\r\n", "\n") for _, _, name, code in parser.extract_macros()}
    finally:
        parser.close()
    assert set(sources) == {"ThisWorkbook.cls", "Sheet1.cls", "Sheet4.cls", "Model_1.bas", "ProgressChart.bas"}
    for name in ["Model_1.bas", "ProgressChart.bas", "Sheet1.cls"]:
        assert sources[name] == (ROOT / "src" / name).read_text()
    ole = olefile.OleFileIO(io.BytesIO(parts["xl/vbaProject.bin"]), raise_defects=olefile.DEFECT_INCORRECT)
    assert not ole.parsing_issues
    assert not any("__SRP" in "/".join(n) for n in ole.listdir())
    project = VBA_Project(ole, "", "PROJECT", "VBA/dir")
    project.parse_project_stream()
    list(project.parse_modules())
    assert len(project.modules) == 5
    assert all(m.textoffset == 0 for m in project.modules)
    assert ole.openstream("VBA/_VBA_PROJECT").read() == bytes.fromhex("cc61ffff000000")
    metadata = ole.openstream("PROJECT").read().decode("cp936")
    assert "Sheet2" not in metadata and "Sheet3" not in metadata


def test_buttons_have_real_public_entry_points(parts):
    code = (ROOT / "src/ProgressChart.bas").read_text()
    entry_points = set(re.findall(r"Public Sub (\w+)\(\)", code))
    bindings = []
    for name in ["xl/drawings/drawing1.xml", "xl/drawings/drawing2.xml"]:
        drawing = ET.fromstring(parts[name])
        ids = [n.get("id") for n in drawing.findall(".//xdr:cNvPr", NS)]
        assert len(ids) == len(set(ids))
        for node in drawing.findall(".//xdr:sp", NS):
            macro = node.get("macro")
            if macro:
                assert macro.startswith("[0]!ProgressChart.")
                routine = macro.split(".")[-1]
                assert routine in entry_points
                bindings.append(routine)
    assert set(bindings) == {"OpenProgressChart", "ImportSelectedMilestones", "GenerateProgressChart", "CopyProgressChart", "RefreshAll", "WidenProgressChart", "NarrowProgressChart", "AddTask", "AddSubtask", "UpdateGantt", "PromoteTasks", "DemoteTasks"}


def test_example_workbook_loads_and_dates_are_real(artifact):
    with warnings.catch_warnings():
        # Legacy extensions are preserved by our XML edits; openpyxl is read-only here.
        warnings.simplefilter("ignore", UserWarning)
        wb = openpyxl.load_workbook(artifact, keep_vba=True)
    assert wb.sheetnames == ["WBS", "Milestones"]
    ws = wb["Milestones"]
    assert wb.active == wb["WBS"]
    assert ws.sheet_properties.codeName == "Sheet4"
    assert ws.freeze_panes == "A14"
    assert ws["C6"].value == "Equal spacing"
    assert ws["C7"].value == 1
    for r, (label, date) in enumerate(EXAMPLES, 14):
        assert ws.cell(r, 2).value == "Yes"
        assert ws.cell(r, 4).value == label
        assert ws.cell(r, 5).value.date() == date
    assert ws.data_validations.dataValidation[0].sqref == "B14:B213"
    assert ws.data_validations.dataValidation[1].sqref == "C6"
    assert ws.data_validations.dataValidation[2].sqref == "C7"
    assert '"Yes,No"' == ws.data_validations.dataValidation[0].formula1
    assert '"Equal spacing,Date proportional"' == ws.data_validations.dataValidation[1].formula1
    assert ws.data_validations.dataValidation[2].formula1 == "0.5"
    assert ws.data_validations.dataValidation[2].formula2 == "2"
    with ZipFile(artifact) as z:
        sheet = ET.fromstring(z.read("xl/worksheets/sheet1.xml"))
        predecessor = sheet.find("s:dataValidations/s:dataValidation[@sqref='O6:O1000']", NS)
        assert predecessor.get("showErrorMessage") == "0"
    assert len(ws.merged_cells.ranges) == 8
    wb.close()


def test_saved_schedule_uses_leaf_rollups_and_calendar_days(artifact):
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(artifact, keep_vba=True, data_only=True)
    ws = wb["WBS"]
    assert ws["B1"].value == "WBS Gantt Planner v5.1"
    tasks = {r: {"level": ws.cell(r, 25).value, "kind": ws.cell(r, 24).value,
                 "start": ws.cell(r, 7).value, "end": ws.cell(r, 8).value,
                 "days": ws.cell(r, 6).value, "progress": ws.cell(r, 14).value}
             for r in range(6, ws.max_row+1) if ws.cell(r, 3).value}
    names = {ws.cell(r, 3).value: r for r in tasks}
    assert len(tasks) == 14
    assert ws.max_row == 19
    for r, task in tasks.items():
        assert 1 <= task["level"] <= 5
        assert ws.cell(r, 7).number_format == ws.cell(r, 8).number_format == "yyyy/m/d"
        assert task["days"] == (task["end"] - task["start"]).days + 1
        if task["kind"] == "P":
            descendants = []
            for n in range(r+1, ws.max_row+1):
                if n not in tasks: continue
                if tasks[n]["level"] <= task["level"]: break
                if tasks[n]["kind"] == "T": descendants.append(tasks[n])
            assert task["start"] == min(t["start"] for t in descendants)
            assert task["end"] == max(t["end"] for t in descendants)
            assert task["progress"] == pytest.approx(sum(t["days"]*t["progress"] for t in descendants)/sum(t["days"] for t in descendants))
        elif ws.cell(r,15).value:
            pred = tasks[names[ws.cell(r,15).value]]
            assert task["start"] == pred["end"] + dt.timedelta(days=1)
    dates = [ws.cell(4,c).value for c in range(26,ws.max_column+1)]
    assert dates[0] == min(t["start"] for t in tasks.values())
    assert dates[-1] == max(t["end"] for t in tasks.values())
    assert all(b-a == dt.timedelta(days=1) for a,b in zip(dates,dates[1:]))
    for c, date in enumerate(dates,26):
        assert ws.cell(5,c).value == ["Mon","Tue","Wed","Thu","Fri","Sat","Sun"][date.weekday()]
        if date.weekday() > 4: assert ws.cell(5,c).fill.fgColor.rgb == "FFF1F5F9"
    assert all(ws.cell(r,c).value is None for r in range(1,4) for c in range(26,ws.max_column+1))
    assert str(ws.data_validations.dataValidation[-1].sqref) == "D2"
    wb.close()


def test_sorted_unique_cells_and_valid_merge_ranges(parts):
    from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries
    for name in ["xl/worksheets/sheet1.xml", "xl/worksheets/sheet4.xml"]:
        sheet = ET.fromstring(parts[name])
        rows = sheet.findall("s:sheetData/s:row", NS)
        numbers = [int(r.get("r")) for r in rows]
        assert numbers == sorted(set(numbers))
        for row in rows:
            refs = [coordinate_to_tuple(c.get("r")) for c in row]
            assert refs == sorted(set(refs))
        occupied = set()
        for merge in sheet.findall("s:mergeCells/s:mergeCell", NS):
            a,b,c,d = range_boundaries(merge.get("ref"))
            cells = {(r,k) for r in range(b,d+1) for k in range(a,c+1)}
            assert not occupied.intersection(cells)
            occupied.update(cells)


def test_chart_geometry_and_labels_match_example_data(parts):
    drawing = ET.fromstring(parts["xl/drawings/drawing2.xml"])
    group = drawing.find("xdr:absoluteAnchor/xdr:grpSp", NS)
    shapes = group.findall("xdr:sp", NS)
    assert len(shapes) == 20
    def rect(sp):
        off = sp.find("xdr:spPr/a:xfrm/a:off", NS)
        ext = sp.find("xdr:spPr/a:xfrm/a:ext", NS)
        return [int(off.get("x")), int(off.get("y")), int(ext.get("cx")), int(ext.get("cy"))]
    axis = shapes[1]
    assert axis.find("xdr:spPr/a:solidFill/a:srgbClr", NS).get("val") == "8B0000"
    ax = rect(axis)
    label_rects = []
    positions = []
    for i, (label,date) in enumerate(EXAMPLES):
        triangle, datebox, labelbox = shapes[2+3*i:5+3*i]
        tr = rect(triangle)
        # The right-arrow shaft is 55% of its bounding height. Align to the
        # visible shaft edge, not the arrowhead's enclosing box.
        assert tr[1]+tr[3]/2 == pytest.approx(ax[1]+ax[3]*(1-.55)/2, abs=1)
        date_rect = rect(datebox)
        assert tr[1] - date_rect[1] - date_rect[3] == pytest.approx(2*12700, abs=1)
        assert triangle.find("xdr:spPr/a:solidFill/a:srgbClr", NS).get("val") == "F2A000"
        assert datebox.find(".//a:t", NS).text == f"{date.month}/{date.day}"
        assert "\n".join(t.text for t in labelbox.findall(".//a:t", NS)) == label
        label_rects.append(rect(labelbox)); positions.append(tr[0])
    for a,b in zip(label_rects,label_rects[1:]):
        assert b[0] - a[0] - a[2] >= 10*12700-1
    gaps = [b-a for a,b in zip(positions,positions[1:])]
    assert max(gaps)-min(gaps) <= 1


@pytest.mark.parametrize("data", [b"", b"a", b"abc" * 3000, b" " * 4096, bytes(range(256)) * 40,
                                  random.Random(8).randbytes(5000), random.Random(4).randbytes(8192)])
def test_vba_compression_round_trip(data):
    assert decompress_stream(compress_vba(data)) == data


def test_cfb_small_large_streams_and_red_black_trees():
    streams = {f"VBA/Module{i}": random.Random(i).randbytes(size)
               for i, size in enumerate([0, 1, 63, 64, 65, 511, 512, 4095, 4096, 4097, 12000])}
    streams.update({"PROJECT": b"metadata", "PROJECTwm": b"names", "VBA/模块": b"unicode"})
    data = cfb_write(streams)
    ole = olefile.OleFileIO(io.BytesIO(data), raise_defects=olefile.DEFECT_INCORRECT)
    assert not ole.parsing_issues
    for name, value in streams.items():
        assert ole.openstream(name).read() == value

    def verify(index):
        if index == FREE:
            return 1
        e = ole.direntries[index]
        if e is None:
            e = ole._load_direntry(index)
        left, right = verify(e.sid_left), verify(e.sid_right)
        assert left == right
        if e.color == 0:
            for child in [e.sid_left, e.sid_right]:
                if child != FREE:
                    assert ole.direntries[child].color == 1
        return left + int(e.color == 1)

    verify(ole.root.sid_child)
    for entry in ole.direntries:
        if entry is not None and entry.entry_type == 1:
            verify(entry.sid_child)
